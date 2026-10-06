import re
from pathlib import Path
from bs4 import BeautifulSoup
import fitz  # PyMuPDF
import ebooklib
from ebooklib import epub
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

def extract_text_from_txt(file_path: Path) -> str:
    """Extracts text from a plain TXT file."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception as e:
        raise ValueError(f"Failed to read TXT file: {str(e)}")

def extract_text_from_pdf(file_path: Path) -> str:
    """Extracts text from a PDF file using PyMuPDF (fitz)."""
    try:
        doc = fitz.open(file_path)
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text())
        doc.close()
        return "\n".join(text_parts).strip()
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")

def extract_text_from_epub(file_path: Path) -> str:
    """Extracts text from an EPUB file using ebooklib and BeautifulSoup."""
    try:
        # Suppress ebooklib warnings about epub3
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            book = epub.read_epub(str(file_path))
            
        text_parts = []
        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), "html.parser")
                # Remove script and style elements
                for script in soup(["script", "style"]):
                    script.decompose()
                text = soup.get_text()
                # Clean up spacing
                lines = (line.strip() for line in text.splitlines())
                chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
                text = "\n".join(chunk for chunk in chunks if chunk)
                if text:
                    text_parts.append(text)
        return "\n\n".join(text_parts).strip()
    except Exception as e:
        raise ValueError(f"Failed to extract text from EPUB: {str(e)}")

def extract_text(file_path: Path) -> str:
    """Generic text extractor based on file extension."""
    suffix = file_path.suffix.lower()
    if suffix == ".txt":
        return extract_text_from_txt(file_path)
    elif suffix == ".pdf":
        return extract_text_from_pdf(file_path)
    elif suffix == ".epub":
        return extract_text_from_epub(file_path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")

def auto_detect_chapter_title(text: str, default_title: str = "Untitled Chapter") -> str:
    """
    Attempts to detect the chapter title from the beginning of the text.
    Looks for lines starting with 'Chapter', 'CHAPTER', 'Prologue', 'Epilogue', etc.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()][:5]
    if not lines:
        return default_title
        
    # Pattern to match Chapter titles
    chapter_pattern = re.compile(
        r'^(chapter\s+\d+|chap\.\s+\d+|prologue|epilogue|act\s+\d+|\d+[\.:\s]+[A-Za-z\s]+)', 
        re.IGNORECASE
    )
    
    for line in lines:
        if chapter_pattern.match(line) or (len(line) < 60 and line.isupper()):
            return line
            
    # Fallback to the first line if it's reasonably short, else default
    first_line = lines[0]
    if len(first_line) < 50:
        return first_line
        
    return default_title

def export_to_txt(dest_path: Path, chapters_data: list[dict]):
    """
    Exports the condensed novel as a plain TXT file.
    chapters_data: list of dicts with 'title' and 'condensed_text' keys.
    """
    try:
        with open(dest_path, "w", encoding="utf-8") as f:
            for chapter in chapters_data:
                title = chapter.get("title", "Untitled Chapter")
                content = chapter.get("condensed_text") or ""
                f.write(f"{title}\n")
                f.write("=" * len(title) + "\n\n")
                f.write(content)
                f.write("\n\n" + "-" * 40 + "\n\n")
    except Exception as e:
        raise IOError(f"Failed to export to TXT: {str(e)}")

def export_to_docx(dest_path: Path, chapters_data: list[dict], book_title: str):
    """
    Exports the condensed novel as a Word document (.docx).
    """
    try:
        doc = DocxDocument()
        doc.add_heading(book_title, level=0)
        
        for chapter in chapters_data:
            title = chapter.get("title", "Untitled Chapter")
            content = chapter.get("condensed_text") or ""
            
            doc.add_heading(title, level=1)
            # Add content paragraph by paragraph to preserve line breaks
            for para in content.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            doc.add_page_break()
            
        doc.save(str(dest_path))
    except Exception as e:
        raise IOError(f"Failed to export to DOCX: {str(e)}")

def export_to_pdf(dest_path: Path, chapters_data: list[dict], book_title: str):
    """
    Exports the condensed novel as a PDF document.
    """
    try:
        doc = SimpleDocTemplate(
            str(dest_path),
            pagesize=letter,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72
        )
        
        styles = getSampleStyleSheet()
        
        # Define clean typography styles
        title_style = ParagraphStyle(
            name="BookTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=28,
            leading=34,
            spaceAfter=30
        )
        
        heading_style = ParagraphStyle(
            name="ChapterHeading",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            spaceBefore=20,
            spaceAfter=15,
            keepWithNext=True
        )
        
        body_style = ParagraphStyle(
            name="BookBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=11,
            leading=16,
            spaceAfter=10
        )
        
        story = []
        
        # Add Book Title
        story.append(Paragraph(book_title, title_style))
        story.append(Spacer(1, 40))
        
        for chapter in chapters_data:
            title = chapter.get("title", "Untitled Chapter")
            content = chapter.get("condensed_text") or ""
            
            # Format and add chapter title
            story.append(Paragraph(title, heading_style))
            
            # Format and add paragraphs
            for para in content.split("\n\n"):
                if para.strip():
                    # Replace newlines inside paragraph with break tags
                    cleaned_para = para.strip().replace("\n", "<br/>")
                    story.append(Paragraph(cleaned_para, body_style))
            
            story.append(Spacer(1, 20))
            
        doc.build(story)
    except Exception as e:
        raise IOError(f"Failed to export to PDF: {str(e)}")
